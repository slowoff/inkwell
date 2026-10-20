# 1.46: importing a Word file or PDF no longer makes you wait in the dialog. Choosing a Word file only lays it out and
# makes page 1 (preview); "Print onto page" closes the dialog at once and the pages appear one by one while you can
# keep writing. Checks: quick dialog, all pages placed in order without overlap, page range, one undo step for the
# whole import, undo while running stops it, leaving the page still finishes it on the right page, Stop button,
# PDF import, nothing left behind.
import asyncio, os, tempfile, time
from playwright.async_api import async_playwright
PORT = os.environ.get('INKWELL_PORT', '8766')
URL = f'http://localhost:{PORT}/'
fails = []
def check(ok, what):
    print(('ok   ' if ok else 'FAIL ') + what)
    if not ok: fails.append(what)

TMP = tempfile.mkdtemp()
def make_docx(n, name):
    from docx import Document
    from docx.enum.text import WD_BREAK
    d = Document()
    for i in range(n):
        d.add_heading(f'Worksheet page {i + 1}', 1)
        for k in range(8): d.add_paragraph(f'Question {k + 1}: explain the idea number {k + 1} on page {i + 1} in your own words.')
        if i < n - 1: d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = os.path.join(TMP, name); d.save(p); return p
def make_pdf(n, name):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    p = os.path.join(TMP, name); c = canvas.Canvas(p, pagesize=A4)
    for i in range(n):
        c.setFont('Helvetica', 14)
        for k in range(10): c.drawString(72, 760 - 30 * k, f'PDF page {i + 1}, line {k + 1}: some printed text to read')
        c.showPage()
    c.save(); return p

IMGS = "[...__inkwell.cur.items.values()].filter(i => i.t === 'i').sort((a, b) => a.y - b.y).map(i => [i.y, i.h, i.doc || 0])"

async def main():
    doc6 = make_docx(6, 'six.docx'); doc12 = make_docx(12, 'twelve.docx'); pdf4 = make_pdf(4, 'Four pages.pdf')
    async with async_playwright() as p:
        b = await p.chromium.launch()
        ctx = await b.new_context(viewport={'width': 1180, 'height': 820})
        await ctx.route('**/*.supabase.co/**', lambda r: r.abort())
        page = await ctx.new_page()
        errs = []
        page.on('pageerror', lambda e: errs.append(str(e)))
        await page.goto(URL); await page.wait_for_timeout(3000)
        await page.wait_for_function('window.__inkwell && window.__inkwell.cur', timeout=15000)

        async def pick(path):
            await page.click('#insertBtn'); await page.wait_for_timeout(200)
            await page.evaluate("() => { document.getElementById('toastText').textContent = ''; }")
            t0 = time.time()
            await page.set_input_files('#importFile', path)
            await page.wait_for_function("!/^(Reading file|Turning)/.test(document.getElementById('impStatus').textContent) && !document.getElementById('impGo').disabled && document.querySelector('#importPreview canvas')", timeout=60000)
            return time.time() - t0
        async def wait_done(timeout=90000):
            await page.wait_for_function("/^(Placed|Import stopped)/.test(document.getElementById('toastText').textContent)", timeout=timeout)

        # --- Word, 6 pages
        await page.evaluate("async () => { await __inkwell.newPage(); }"); await page.wait_for_timeout(400)
        pid = await page.evaluate("__inkwell.cur.id")
        dt = await pick(doc6)
        n_prev = await page.evaluate("document.querySelectorAll('#importPreview canvas').length")
        cnt = await page.evaluate("document.getElementById('impCount').textContent")
        hint = await page.evaluate("document.getElementById('impStatus').textContent")
        check('Not signed in' in hint, f'signed out: Word is drawn by Inkwell and the dialog says so ({hint!r})')
        check(n_prev == 1, f'the dialog only makes page 1 for the preview ({n_prev} pictures, {dt:.1f} s)')
        check('6' in cnt, f'page count shown before importing: {cnt!r}')
        await page.click('#impGo')
        await page.wait_for_timeout(150)
        check(await page.evaluate("document.getElementById('importDlg').hidden"), 'the dialog closes right away after Print onto page')
        check(await page.evaluate(f"({IMGS}).length") < 6, 'pages are still being made after the dialog closed (not all at once)')
        # write while it runs
        await page.evaluate("() => __inkwell.setTool('pen')")
        st = await page.evaluate("(() => { const r = document.getElementById('stage').getBoundingClientRect(); return [r.left + r.width / 2, r.top + r.height / 2]; })()")
        await page.mouse.move(st[0], st[1]); await page.mouse.down()
        for k in range(1, 15): await page.mouse.move(st[0] + 8 * k, st[1] + 3 * k)
        await page.mouse.up()
        await wait_done()
        imgs = await page.evaluate(IMGS)
        check(len(imgs) == 6, f'all 6 Word pages placed ({len(imgs)})')
        check(all(x[2] == 1 for x in imgs), 'placed pages are document pages (doc:1)')
        check(all(imgs[i][0] + imgs[i][1] <= imgs[i + 1][0] + 0.5 for i in range(len(imgs) - 1)), 'pages are stacked in order without overlapping')
        strokes = await page.evaluate("[...__inkwell.cur.items.values()].filter(i => i.t === 's').length")
        check(strokes == 1, f'a stroke written during the import is kept ({strokes})')
        title = await page.evaluate("document.getElementById('title').value")
        check(title == 'six', f'empty page takes the file name ({title!r})')
        check(await page.evaluate("!document.querySelector('.docx-render')"), 'the hidden Word layout is removed afterwards')
        # one undo step: the stroke was written after the first page, so undo the stroke, then the import
        await page.evaluate("() => { document.getElementById('undo').click(); document.getElementById('undo').click(); }"); await page.wait_for_timeout(200)
        left = await page.evaluate(f"({IMGS}).length")
        check(left == 0, f'one undo removes the whole import ({left} pages left)')
        await page.evaluate("() => document.getElementById('redo').click()"); await page.wait_for_timeout(200)
        check(await page.evaluate(f"({IMGS}).length") == 6, 'redo brings all pages back')
        await page.evaluate("() => __inkwell.flushSaves()"); await page.wait_for_timeout(300)

        # --- second import onto the same page goes below the first
        y_end = imgs[-1][0] + imgs[-1][1]
        await pick(doc6)
        await page.fill('#impPages', '2-3')
        await page.click('#impGo'); await wait_done()
        imgs2 = await page.evaluate(IMGS)
        check(len(imgs2) == 8, f'page range 2-3 places 2 pages ({len(imgs2) - 6})')
        check(min(x[0] for x in imgs2[6:]) >= y_end, 'a second import goes below the first')

        # --- undo while running stops it
        await page.evaluate("async () => { await __inkwell.newPage(); }"); await page.wait_for_timeout(400)
        await pick(doc12)
        await page.click('#impGo')
        await page.wait_for_function(f"({IMGS}).length >= 1", timeout=30000)
        await page.evaluate("() => document.getElementById('undo').click()")
        await wait_done()
        await page.wait_for_timeout(600)
        n = await page.evaluate(f"({IMGS}).length")
        txt = await page.evaluate("document.getElementById('toastText').textContent")
        check(n == 0 and txt.startswith('Import stopped'), f'undo while importing stops it ({n} pages, toast {txt!r})')

        # --- Stop button
        await pick(doc12)
        await page.click('#impGo')
        await page.wait_for_function(f"({IMGS}).length >= 2", timeout=30000)
        await page.wait_for_function("!document.getElementById('toastBtn').hidden && document.getElementById('toastBtn').textContent === 'Stop'", timeout=20000)
        await page.click('#toastBtn')
        await wait_done(); await page.wait_for_timeout(500)
        n = await page.evaluate(f"({IMGS}).length")
        check(2 <= n < 12, f'Stop ends the import and keeps what is placed ({n} pages)')

        # --- leave the page while it runs
        await page.evaluate("async () => { await __inkwell.newPage(); }"); await page.wait_for_timeout(400)
        pa = await page.evaluate("__inkwell.cur.id")
        await pick(doc12)
        await page.click('#impGo')
        await page.wait_for_function(f"({IMGS}).length >= 1", timeout=30000)
        await page.evaluate("async () => { await __inkwell.newPage(); }")
        pb = await page.evaluate("__inkwell.cur.id")
        await wait_done()
        on_b = await page.evaluate(f"({IMGS}).length")
        check(on_b == 0, f'the other page gets nothing ({on_b})')
        await page.evaluate(f"async () => {{ await __inkwell.openPage('{pa}'); }}"); await page.wait_for_timeout(500)
        ia = await page.evaluate(IMGS)
        check(len(ia) == 12, f'leaving the page: the rest still goes onto the page the import started on ({len(ia)})')
        check(all(ia[i][0] + ia[i][1] <= ia[i + 1][0] + 0.5 for i in range(len(ia) - 1)), '... in order without overlapping')
        await page.reload(); await page.wait_for_timeout(3000)
        await page.evaluate(f"async () => {{ await __inkwell.openPage('{pa}'); }}"); await page.wait_for_timeout(500)
        check(await page.evaluate(f"({IMGS}).length") == 12, 'all pages are still there after a reload')

        # --- PDF
        await page.evaluate("async () => { await __inkwell.newPage(); }"); await page.wait_for_timeout(400)
        await pick(pdf4)
        await page.click('#impGo')
        await page.wait_for_timeout(100)
        check(await page.evaluate("document.getElementById('importDlg').hidden"), 'PDF: dialog closes right away')
        await wait_done()
        ip = await page.evaluate(IMGS)
        check(len(ip) == 4, f'PDF: all 4 pages placed ({len(ip)})')
        check(await page.evaluate("document.getElementById('title').value") == 'Four pages', 'PDF: page takes the file name')
        await page.wait_for_timeout(500)
        np = await page.evaluate("[...__inkwell.docs.values()].filter(d => d.kind === 'ptxt' && !d.deleted).length")
        check(np >= 4, f'PDF text is searchable right away ({np} pages of text)')

        # --- signed in: Word goes through Google Docs and comes back as a PDF
        calls = []
        pdf_bytes = open(make_pdf(3, 'conv.pdf'), 'rb').read()
        async def up(route):
            calls.append(('upload', route.request.post_data_buffer or b''))
            await route.fulfill(status=200, content_type='application/json', body='{"id":"gdoc1"}')
        async def ex(route):
            if route.request.method == 'DELETE':
                calls.append(('delete', b'')); await route.fulfill(status=204, body='')
            else:
                calls.append(('export', route.request.url.encode())); await route.fulfill(status=200, content_type='application/pdf', body=pdf_bytes)
        await page.route('**/upload/drive/v3/files**', up)
        await page.route('**/drive/v3/files/gdoc1**', ex)
        await page.evaluate("() => { __inkwell.Sync.user = {id: 'u1', email: 't@example.com'}; __inkwell.GD.token = async () => ({at: 'tok', folder: 'fold', exp: Date.now() + 3e6}); }")
        await page.evaluate("async () => { await __inkwell.newPage(); }"); await page.wait_for_timeout(400)
        await pick(doc6)
        kinds = [c[0] for c in calls]
        check(kinds[:2] == ['upload', 'export'], f'signed in: the Word file is converted by Google ({kinds})')
        check(b'application/vnd.google-apps.document' in calls[0][1] and b'Worksheet page' not in calls[0][1][:0] and len(calls[0][1]) > 5000, 'uploaded as a Google Doc with the Word file in it')
        check(b'mimeType=application%2Fpdf' in calls[1][1], 'exported as PDF')
        await page.wait_for_timeout(300)
        check('delete' in [c[0] for c in calls], 'the temporary Google Doc is deleted again')
        cnt = await page.evaluate("document.getElementById('impCount').textContent")
        st = await page.evaluate("document.getElementById('impStatus').textContent")
        check(cnt == '(3 pages)' and st == '', f'the dialog shows the PDF pages, no warning ({cnt!r}, {st!r})')
        await page.click('#impGo'); await wait_done()
        ig = await page.evaluate(IMGS)
        check(len(ig) == 3, f'converted Word file placed like a PDF ({len(ig)} pages)')
        check(await page.evaluate("document.getElementById('title').value") == 'six', 'page named after the Word file')
        # Google fails: falls back to Inkwell's own drawing with a warning
        await page.unroute('**/upload/drive/v3/files**')
        await page.route('**/upload/drive/v3/files**', lambda r: r.fulfill(status=500, body='{}'))
        await pick(doc6)
        st = await page.evaluate("document.getElementById('impStatus').textContent")
        check('could not convert' in st and 'about' in await page.evaluate("document.getElementById('impCount').textContent"), f'conversion fails: Inkwell draws it itself and says so ({st!r})')
        await page.click('#importDlg [data-close]'); await page.wait_for_timeout(300)
        await page.evaluate("() => { __inkwell.Sync.user = null; }")

        # --- closing the dialog after picking a Word file leaves nothing behind
        await pick(doc6)
        await page.click('#importDlg [data-close]'); await page.wait_for_timeout(300)
        check(await page.evaluate("!document.querySelector('.docx-render')"), 'closing the dialog drops the opened Word file')

        check(not errs, f'no page errors {errs}')
        await b.close()
    print('ALL OK' if not fails else f'{len(fails)} FAILED')
    raise SystemExit(1 if fails else 0)

asyncio.run(main())

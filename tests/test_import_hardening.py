# 1.67.1 (security check 2026-10-27): imported files cannot run code inside Inkwell.
# Word: an embedded web page (altChunk) is not shown (docx-preview put it in an iframe without a sandbox).
# PDF: opened with isEvalSupported: false (pdf.js 3.x font code built with new Function, CVE-2024-4367); import still works.
import asyncio, os, tempfile
from playwright.async_api import async_playwright
PORT = os.environ.get('INKWELL_PORT', '8766')
URL = f'http://localhost:{PORT}/'
fails = []
def check(ok, what):
    print(('ok   ' if ok else 'FAIL ') + what)
    if not ok: fails.append(what)

TMP = tempfile.mkdtemp()
def make_docx_altchunk(name):
    from docx import Document
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.oxml import parse_xml
    d = Document()
    d.add_paragraph('Normal text before.')
    html = b'<html><body><p>ALTCHUNK web page text</p></body></html>'
    part = Part(PackURI('/word/chunk1.html'), 'text/html', html, d.part.package)
    rid = d.part.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk')
    d.element.body.append(parse_xml(f'<w:altChunk xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid}"/>'))
    d.add_paragraph('Normal text after.')
    p = os.path.join(TMP, name); d.save(p); return p
def make_pdf(name):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    p = os.path.join(TMP, name); c = canvas.Canvas(p, pagesize=A4)
    c.setFont('Helvetica', 14)
    for k in range(10): c.drawString(72, 760 - 30 * k, f'Line {k + 1}: printed text')
    c.showPage(); c.save(); return p

async def main():
    docx = make_docx_altchunk('chunk.docx'); pdf = make_pdf('one.pdf')
    async with async_playwright() as p:
        b = await p.chromium.launch()
        ctx = await b.new_context(viewport={'width': 1180, 'height': 820})
        await ctx.route('**/*.supabase.co/**', lambda r: r.abort())
        page = await ctx.new_page(); errs = []
        page.on('pageerror', lambda e: errs.append(str(e)))
        await page.goto(URL); await page.wait_for_timeout(3000)
        await page.wait_for_function('!!window.__inkwellBooted', timeout=15000)
        await page.evaluate("""async () => { await __inkwell.newPage();
          window.__frames = 0; new MutationObserver(ms => { for(const m of ms) for(const n of m.addedNodes){ const f = n.nodeType === 1 && (n.tagName === 'IFRAME' ? n : n.querySelector && n.querySelector('iframe'));
            // html2canvas makes its own hidden copy of the page in an iframe: that one is not counted
            if(f && !f.classList.contains('html2canvas-container')){ window.__frames++; window.__frameInfo = (window.__frameInfo || []).concat([f.outerHTML.slice(0, 200)]); } } })
            .observe(document.documentElement, {childList: true, subtree: true});
}""")

        # Word file with an embedded web page
        await page.click('#insertBtn'); await page.wait_for_timeout(200)
        await page.set_input_files('#importFile', docx)
        await page.wait_for_function("!/^(Reading file|Turning)/.test(document.getElementById('impStatus').textContent) && document.querySelector('#importPreview canvas')", timeout=60000)
        n = await page.evaluate("window.__frames")
        check(n == 0, f'Word: no iframe is made for the embedded web page ({n} {await page.evaluate("window.__frameInfo")})')
        await page.click('#impGo')
        await page.wait_for_function("/^Placed/.test(document.getElementById('toastText').textContent)", timeout=60000)
        pics = await page.evaluate("[...__inkwell.cur.items.values()].filter(i => i.t === 'i').length")
        check(pics >= 1, f'Word: the normal text is still imported ({pics} page)')

        # PDF
        await page.click('#insertBtn'); await page.wait_for_timeout(200)
        await page.evaluate("() => { document.getElementById('toastText').textContent = ''; }")
        await page.set_input_files('#importFile', pdf)
        await page.wait_for_function("!/^(Reading file|Turning)/.test(document.getElementById('impStatus').textContent) && document.querySelector('#importPreview canvas')", timeout=60000)
        src = await page.evaluate("String(window.pdfOpen)")
        check('isEvalSupported: false' in src, 'PDF: opened with isEvalSupported: false')
        ink = await page.evaluate("""(() => { const c = document.querySelector('#importPreview canvas'); const d = c.getContext('2d').getImageData(0, 0, c.width, c.height).data; let k = 0; for(let i = 0; i < d.length; i += 4) if(d[i] < 200) k++; return k; })()""")
        check(ink > 50, f'PDF: the text is still drawn ({ink} dark pixels)')
        check(not errs, f'no page errors {errs[:2]}')
        await b.close()

asyncio.run(main())
print('ALL OK' if not fails else f'{len(fails)} FAILED')
raise SystemExit(1 if fails else 0)
